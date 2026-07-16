# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.shared import Pt

path = Path(__file__).resolve().parents[1] / "Documentacion" / "EduplanerNocturna" / "MANUAL_PREMIUM_CONFIGURACION_INICIAL_EDUPLANER_NOCTURNA.docx"
doc = Document(str(path))
texts = "\n".join(p.text for p in doc.paragraphs)
marker = "Anexo A — Validación instalación limpia local"
if marker not in texts:
    doc.add_page_break()
    doc.add_heading(marker, level=1)
    bullets = [
        "Fecha: 2026-07-16. BD local eduplaner_nocturna_local (PostgreSQL 18). App http://localhost:5172.",
        "Migraciones: se corrigió creación de shifts antes de time_slots y shift_id en student_assignments/groups para instalación limpia.",
        "Orden validado: SuperAdmin escuela+admin → jornada Noche → SaveCatalog → trimestres → usuarios → /SaveAssignments → ScheduleConfiguration nocturna → Schedule/ByTeacher → PrematriculationPeriod → StudentAssignment.",
        "Evidencia: 7 time_slots nocturnos (18:00–22:45), 4 imparticiones, 1 schedule_entry, 1 matrícula.",
        "Importante: StudentAssignment no crea automáticamente student_subject_assignments; el flujo de notas por materia/modular requiere SSA vía prematricula modular.",
        "Credenciales prueba: superadmin@schoolmanager.com / Admin123! ; admin.prueba@celosam.local / Admin123! ; docente.prueba@celosam.local / Teacher123!.",
        "Entregables: CHECKLIST_FINAL_IMPLEMENTACION.md, INFORME_VALIDACION_FINAL_LOCAL.md, tres matrices XLSX en Documentacion/EduplanerNocturna/.",
    ]
    for b in bullets:
        p = doc.add_paragraph(b, style="List Bullet")
        for r in p.runs:
            r.font.size = Pt(11)
    doc.save(str(path))
    print("Annex appended")
else:
    print("Annex already present")
print("size", path.stat().st_size)
