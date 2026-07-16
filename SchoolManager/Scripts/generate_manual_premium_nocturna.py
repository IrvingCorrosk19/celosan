#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera MANUAL_PREMIUM_CONFIGURACION_INICIAL_EDUPLANER_NOCTURNA.docx"""
from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor, Cm

OUT_DIR = Path(r"C:\Proyectos\EduplanerNoche\SchoolManager\Documentacion\EduplanerNocturna")
OUT = OUT_DIR / "MANUAL_PREMIUM_CONFIGURACION_INICIAL_EDUPLANER_NOCTURNA.docx"

NAVY = RGBColor(0x0B, 0x2C, 0x4A)
ACCENT = RGBColor(0x1A, 0x5F, 0x9E)
GRAY = RGBColor(0x4A, 0x4A, 0x4A)
GREEN = RGBColor(0x1B, 0x7A, 0x3D)
ORANGE = RGBColor(0xC0, 0x56, 0x00)
RED = RGBColor(0xA1, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def shade(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run(run, size=11, bold=False, color=None, italic=False, name="Calibri"):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rPr.append(rFonts)
    if color:
        run.font.color.rgb = color


def add_hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A5F9E")
    pBdr.append(bottom)
    pPr.append(pBdr)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        set_run(r, size={0: 22, 1: 16, 2: 13, 3: 11}.get(level, 11), bold=True, color=NAVY if level <= 1 else ACCENT)
    return h


def para(doc, text, size=11, bold=False, color=GRAY, align=None, space_after=8, italic=False):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    for r in p.runs:
        set_run(r, size=10, color=GRAY)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(text, style="List Number")
    for r in p.runs:
        set_run(r, size=10, color=GRAY)
    return p


def callout(doc, title, body, fill="FFF4E5", border="C05600"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade(cell, fill)
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    set_run(r, size=10, bold=True, color=ORANGE if border == "C05600" else NAVY)
    r2 = p.add_run(body)
    set_run(r2, size=10, color=GRAY)
    doc.add_paragraph()


def table(doc, headers, rows, header_fill="0B2C4A"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        shade(c, header_fill)
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        set_run(r, size=9, bold=True, color=WHITE)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.rows[i + 1].cells[j]
            if i % 2 == 1:
                shade(c, "F3F6FA")
            c.text = ""
            r = c.paragraphs[0].add_run(str(val))
            set_run(r, size=9, color=GRAY)
    doc.add_paragraph()
    return t


def marker(doc, caption):
    callout(
        doc,
        "📷 MARCADOR DE CAPTURA DE PANTALLA",
        f"Insertar aquí captura real de: {caption}. Ejecutar la aplicación localmente o en el entorno de la institución y pegar la imagen con etiquetas de numeración.",
        fill="E8F1FB",
        border="1A5F9E",
    )


def procedure(doc, **kw):
    heading(doc, kw["title"], 2)
    para(doc, f"Objetivo: {kw['objetivo']}", bold=True, color=NAVY)
    para(doc, kw["descripcion"])
    para(doc, "Dependencias", bold=True, color=ACCENT, size=11)
    for d in kw.get("deps", []):
        bullet(doc, d)
    para(doc, f"Ruta exacta: {kw['ruta']}", bold=True)
    para(doc, f"Roles autorizados: {kw['roles']}")
    para(doc, f"Pantalla: {kw['pantalla']}")
    marker(doc, kw["pantalla"])
    if kw.get("campos"):
        para(doc, "Campos (obligatorios marcados *)", bold=True, color=ACCENT)
        table(doc, ["Campo", "Obligatorio", "Descripción"], kw["campos"])
    para(doc, "Paso a paso", bold=True, color=ACCENT)
    for i, step in enumerate(kw["pasos"], 1):
        numbered(doc, step)
    if kw.get("ejemplo"):
        para(doc, f"Ejemplo: {kw['ejemplo']}", italic=True)
    para(doc, f"Resultado esperado: {kw['resultado']}", bold=True, color=GREEN)
    para(doc, f"Validación: {kw['validacion']}")
    if kw.get("errores"):
        para(doc, "Errores frecuentes y solución", bold=True, color=ORANGE)
        table(doc, ["Error / síntomas", "Causa probable", "Solución"], kw["errores"])
    para(doc, f"Impacto: {kw['impacto']}")
    if kw.get("buenas"):
        para(doc, "Buenas prácticas", bold=True, color=ACCENT)
        for b in kw["buenas"]:
            bullet(doc, b)
    if kw.get("checklist"):
        para(doc, "Checklist", bold=True, color=ACCENT)
        for c in kw["checklist"]:
            bullet(doc, f"☐ {c}")
    add_hr(doc)


def footer_header(doc):
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run("Eduplaner Nocturna · Manual Premium de Configuración Inicial · Confidencial")
        set_run(r, size=8, color=ACCENT, italic=True)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # page number field
        run = fp.add_run("Página ")
        set_run(run, size=8, color=GRAY)
        fldChar1 = OxmlElement("w:fldChar")
        fldChar1.set(qn("w:fldCharType"), "begin")
        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")
        instrText.text = "PAGE"
        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "end")
        run2 = fp.add_run()
        run2._r.append(fldChar1)
        run2._r.append(instrText)
        run2._r.append(fldChar2)
        r3 = fp.add_run(" | Documentación Oficial | julio 2026")
        set_run(r3, size=8, color=GRAY)


def cover(doc):
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("EDUPLANER NOCTURNA")
    set_run(r, size=14, bold=True, color=ACCENT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MANUAL OFICIAL PREMIUM")
    set_run(r, size=28, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Configuración Inicial del Sistema")
    set_run(r, size=18, color=ACCENT)

    add_hr(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Guía de implementación institucional para escuelas de educación\n"
        "laboral / nocturna (CELOSAM y equivalentes)\n\n"
        "Calidad de consultoría enterprise · ASP.NET Core · PostgreSQL"
    )
    set_run(r, size=11, color=GRAY)

    doc.add_paragraph()
    table(
        doc,
        ["Atributo", "Valor"],
        [
            ["Producto", "SchoolManager / Eduplaner Nocturna"],
            ["Versión documento", "1.0"],
            ["Fecha", "15 de julio de 2026"],
            ["Código base", r"C:\Proyectos\EduplanerNoche\SchoolManager"],
            ["Clasificación", "Confidencial — Uso institucional"],
            ["Audiencia", "Administradores, secretaría, dirección, implementadores"],
            ["Documento hermano", "ANALISIS_COMPLETO_EDUPLANER_NOCTURNA.md"],
        ],
        header_fill="0B2C4A",
    )

    callout(
        doc,
        "⚠ Principio de veracidad",
        "Este manual solo documenta pantallas, rutas y comportamientos presentes en el código fuente. "
        "Lo que no existe se declara explícitamente. No se inventan botones, menús ni procesos.",
        fill="FDECEA",
        border="A11A1A",
    )
    doc.add_page_break()


def toc_placeholder(doc):
    heading(doc, "Tabla de contenido", 1)
    para(
        doc,
        "En Microsoft Word: referencias → Tabla de contenido → Actualizar campos (clic derecho sobre este párrafo tras abrir el archivo). "
        "Los encabezados del documento están estructurados en niveles Heading 1–3 para TOC automática.",
    )
    for item in [
        "1. Control de versiones e índices",
        "2. Resumen ejecutivo",
        "3. Alcance, supuestos y preparación",
        "4. Orden canónico de configuración (dependencias reales)",
        "5. Roles, permisos y seguridad",
        "6. Preparación inicial y Tenant/Escuela",
        "7. Año lectivo",
        "8. Trimestres",
        "9. Jornada nocturna y parámetros",
        "10. Catálogo académico (grados, grupos, materias, áreas, especialidades)",
        "11. Impartición (SubjectAssignment)",
        "12. Profesores y TeacherAssignment",
        "13. Estudiantes, acudientes y matrícula",
        "14. Prematrícula y pagos",
        "15. Horarios (configuración, bloques, carga por docente)",
        "16. Evaluación, actividades y asistencia",
        "17. Reportes, CELOSAM, carnets y portales",
        "18. Validación final y certificación operativa",
        "19. Matrices, diagramas y Anexos",
        "20. Glosario y limitaciones conocidas",
    ]:
        bullet(doc, item)
    doc.add_page_break()


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    footer_header(doc)
    cover(doc)
    toc_placeholder(doc)

    # 1 Versions
    heading(doc, "1. Control de versiones", 1)
    table(
        doc,
        ["Versión", "Fecha", "Autor", "Cambio"],
        [
            ["1.0", "2026-07-15", "Arquitectura + Análisis código", "Emisión inicial Manual Premium Configuración Inicial"],
        ],
    )
    heading(doc, "1.1 Índice de tablas", 2)
    bullet(doc, "Tablas de roles, dependencias, procedimientos y checklists insertadas en cada capítulo.")
    heading(doc, "1.2 Índice de imágenes / capturas", 2)
    para(doc, "Las capturas se marcan con recuadros azules «MARCADOR DE CAPTURA». Insertar pantallas reales del entorno de la institución.")
    doc.add_page_break()

    # 2 Executive
    heading(doc, "2. Resumen ejecutivo", 1)
    para(
        doc,
        "Eduplaner Nocturna permite operar una institución educativa laboral nocturna (modelo CELOSAM) "
        "sobre SchoolManager: catálogo académico, matrícula/prematrícula, horarios por docente, evaluación, "
        "asistencia, pagos, carnets y módulos de malla curricular.",
    )
    para(doc, "Hallazgos clave", bold=True, color=NAVY)
    bullet(doc, "La configuración de jornada genera SOLO bloques de Noche.")
    bullet(doc, "EnrollmentType por defecto = Nocturno.")
    bullet(doc, "No hay pantalla Index dedicada de Año Académico; se crea al bootstrap de escuela.")
    bullet(doc, "La matrícula operativa es StudentAssignment (no existe entidad 'Matricula').")
    bullet(doc, "Horarios dependen de TeacherAssignment + TimeSlot + AcademicYear.")
    bullet(doc, "Flags: NocturnalAdvancedEnrollment y NocturnalModularEnrollment activos en appsettings.")

    heading(doc, "2.1 Objetivos de este manual", 2)
    for x in [
        "Implementar desde cero una institución nocturna",
        "Capacitar administradores y secretaría",
        "Auditar configuración",
        "Certificar go-live",
        "Soportar operaciones post-implementación",
    ]:
        bullet(doc, x)
    doc.add_page_break()

    # 3 Prep
    heading(doc, "3. Alcance, supuestos y preparación inicial", 1)
    heading(doc, "3.1 Requisitos técnicos", 2)
    bullet(doc, "Servidor/aplicación ASP.NET Core 8 desplegada (local o Render).")
    bullet(doc, "PostgreSQL accesible con migraciones aplicadas.")
    bullet(doc, "Usuario superadmin inicial disponible.")
    bullet(doc, "Navegador moderno (Chrome/Edge).")
    bullet(doc, "Zona horaria configurada: America/Panama (appsettings DateTime.DisplayTimeZoneId).")

    heading(doc, "3.2 Checklist de preparación", 2)
    for c in [
        "Acceso URL del sistema",
        "Credenciales SuperAdmin",
        "Datos oficiales institución (nombre, dirección, logo)",
        "Listado de grados/grupos nocturnos",
        "Listado docentes y cédulas/correos",
        "Malla de materias por especialidad",
        "Calendario de trimestres 2026",
        "Horario oficial de bloques nocturnos (ej. 18:10–23:10)",
    ]:
        bullet(doc, f"☐ {c}")
    callout(
        doc,
        "Buenas prácticas",
        "Complete primero el catálogo y asignaciones docentes ANTES de abrir prematrícula a acudientes. "
        "Abrir el portal sin SubjectAssignments nocturnos produce grados vacíos.",
    )
    doc.add_page_break()

    # 4 Order
    heading(doc, "4. Orden canónico de configuración (dependencias reales)", 1)
    para(doc, "Este orden proviene de FKs y validaciones del código — no es arbitrario.", bold=True)
    table(
        doc,
        ["Paso", "Entidad", "Ruta principal", "Bloquea si falta"],
        [
            ["0", "Escuela + Admin", "/SuperAdmin/CreateSchoolWithAdmin", "Todo el tenant"],
            ["1", "Año académico", "Auto EnsureDefaultAcademicYear", "Horarios, matrícula, períodos"],
            ["2", "Jornada Noche", "/AcademicCatalog (Jornadas)", "Grupos nocturnos / TimeSlots"],
            ["3", "Catálogo GL/Grupo/Materia/Área/Esp.", "/AcademicCatalog/Index", "SubjectAssignment"],
            ["4", "Trimestres", "/AcademicCatalog (Trimestres)", "Actividades / gradebook"],
            ["5", "Impartición", "/SubjectAssignment/Index", "Docentes + prematrícula"],
            ["6", "Usuarios", "/User/Index", "Asignaciones humanas"],
            ["7", "Docente→Impartición", "/TeacherAssignment/Index", "Horario + notas"],
            ["8", "Bloques Noche", "/ScheduleConfiguration + /TimeSlot/Manage", "ScheduleEntry"],
            ["9", "Horario por docente", "/Schedule/ByTeacher", "Vista estudiante"],
            ["10", "Período prematrícula", "/PrematriculationPeriod", "Flujo portal acudiente"],
            ["11", "Matrícula (StudentAssignment)", "/StudentAssignment o ConfirmMatriculation", "Notas/asistencia"],
            ["12", "Ops evaluación/asistencia", "Gradebook / Attendance", "Operación diaria"],
        ],
    )

    heading(doc, "4.1 Diagrama textual de dependencias", 2)
    para(
        doc,
        "School → AcademicYear + Shift(Noche) → Catalog(GL, Group, Subject, Area, Specialty, Trimester) → "
        "SubjectAssignment → Users(Teacher/Student/Acudiente) → TeacherAssignment → "
        "ScheduleConfiguration/TimeSlots → ScheduleEntry → PrematriculationPeriod → Prematriculation/Pago → "
        "StudentAssignment → Activities/Gradebook/Attendance/Reportes/Portales",
        size=10,
    )
    doc.add_page_break()

    # 5 Roles
    heading(doc, "5. Roles, permisos y seguridad", 1)
    para(doc, "Los roles viven en users.role (string). Cookie auth + policies.")
    table(
        doc,
        ["Rol", "Propósito principal", "Acceso típico configuración"],
        [
            ["superadmin", "Crear escuelas, gobierno multi-tenant", "CreateSchool, settings globales"],
            ["admin", "Administración escolar completa", "Catálogo, asignaciones, horarios, usuarios"],
            ["director", "Supervisión académica", "Horarios, reportes, planes, estudiantes"],
            ["secretaria", "Operación secretaría", "Catálogo (Authorize), carnets, CELOSAM"],
            ["teacher / docente", "Enseñanza", "Gradebook, ByTeacher propio, asistencia"],
            ["estudiante / student", "Portal alumno", "Prematrícula, horario, reportes"],
            ["acudiente / parent", "Portal familia", "Prematrícula hijos, pagos, ParentAcademic"],
            ["contable / contabilidad", "Caja y pagos", "Payment"],
            ["clubparentsadmin", "Club de padres", "/ClubParents/Students"],
            ["inspector / qlservices", "Soporte / carnets externos", "Según controllers"],
        ],
    )
    callout(
        doc,
        "Seguridad",
        "Todo usuario escolar nace con SchoolId del admin creador. SuperAdmin tiene SchoolId null. "
        "No comparta la contraseña del SuperAdmin. Use /ChangePassword tras el primer acceso.",
        fill="E8F1FB",
        border="1A5F9E",
    )
    doc.add_page_break()

    # Procedures core
    procedure(
        doc,
        title="6. Crear institución (School) y administrador",
        objetivo="Crear el tenant institucional y su primer administrador.",
        descripcion="Solo el rol superadmin puede crear escuelas. Al crear, el sistema asegura año académico por defecto y puede generar time slots iniciales (legacy mañana) que luego deben reemplazarse por configuración nocturna.",
        deps=["Usuario superadmin autenticado", "Datos legales de la institución"],
        ruta="/SuperAdmin/CreateSchoolWithAdmin",
        roles="superadmin",
        pantalla="SuperAdmin · Crear escuela con administrador",
        campos=[
            ["Nombre escuela", "*", "Razón social / nombre oficial"],
            ["Dirección / teléfono", "Recomendado", "Datos de contacto"],
            ["Datos admin (nombre, email, password)", "*", "Primer usuario admin del tenant"],
        ],
        pasos=[
            "Inicie sesión como superadmin.",
            "Navegue a /SuperAdmin/CreateSchoolWithAdmin.",
            "Complete datos de escuela y del administrador.",
            "Guarde y verifique ListSchools.",
            "Cierre sesión superadmin e inicie con el admin de la escuela.",
        ],
        ejemplo="CELOSAM San Miguelito + admin.celosam@institucion.edu",
        resultado="Escuela IsActive=true; admin con SchoolId; AcademicYear por defecto creado vía EnsureDefaultAcademicYearForSchoolAsync.",
        validacion="Login admin → Dashboard /Home/Index sin AccessDenied.",
        errores=[
            ["Email duplicado", "Unique users.email", "Usar otro correo"],
            ["No aparece año en horarios", "Año inactivo / duplicados", "Verificar academic_years IsActive"],
        ],
        impacto="Base de todo el multi-tenant.",
        buenas=["Un solo admin operativo inicial", "Logo institucional cargar luego", "Documentar credenciales en gestor seguro"],
        checklist=["Escuela visible", "Admin puede entrar", "SchoolId no nulo en admin"],
    )

    procedure(
        doc,
        title="7. Año lectivo (AcademicYear)",
        objetivo="Disponer de un año académico activo único para horarios y matrículas.",
        descripcion="NO existe /AcademicYear/Index. El año se crea automáticamente. En producción CELOSAM deben evitarse múltiples 2026 activos (tuvo historial de duplicados).",
        deps=["School existente"],
        ruta="Servicio AcademicYearService.EnsureDefaultAcademicYearForSchoolAsync · utilitario /Prematriculation/ApplyAcademicYearChangesPage",
        roles="Sistema / admin (utilitarios)",
        pantalla="Sin UI Index dedicada — verificar en dropdown de /Schedule/ByTeacher",
        campos=[["Name", "*", "Ej. 2026"], ["StartDate/EndDate", "*", "Rango del año"], ["IsActive", "*", "Solo uno recomendado"]],
        pasos=[
            "Tras crear escuela, abra /Schedule/ByTeacher y confirme que aparece un año.",
            "Si hay varios años activos del mismo nombre, deje uno canónico IsActive=true y desactive el resto (SQL/admin DB o utilidades).",
            "No cree años duplicados sin necesidad.",
        ],
        ejemplo="2026 activo desde 2026-01-01",
        resultado="Un AcademicYear activo por escuela para el ciclo vigente.",
        validacion="Dropdown Año académico en ByTeacher lista el año y carga tabla.",
        errores=[["No hay años académicos", "Bootstrap falló", "Ejecutar EnsureDefault / ApplyAcademicYearChanges"]],
        impacto="ScheduleEntry, PrematriculationPeriod, StudentAssignment y scores dependen del año.",
        buenas=["Un año canónico", "Nombre corto estable (2026)"],
        checklist=["Un IsActive=true", "Visible en ByTeacher"],
    )

    procedure(
        doc,
        title="8. Configurar trimestres",
        objetivo="Definir el calendario de evaluación del año escolar.",
        descripcion="En Catálogo Académico, pestaña Trimestres. Permite 3 o 4 trimestres, fechas, activar/desactivar y editar.",
        deps=["School", "Recomendable AcademicYear existente"],
        ruta="/AcademicCatalog/Index → pestaña Trimestres",
        roles="admin (menú); Authorize también secretaria, director",
        pantalla="Configuración del Año Escolar y Trimestres",
        campos=[
            ["Cantidad de Trimestres", "*", "3 (recomendado) o 4"],
            ["Trimestre N - Inicio/Fin", "*", "Fechas no solapadas"],
            ["Estado activo", "*", "Activar el trimestre vigente"],
        ],
        pasos=[
            "Ir a Administración → Catálogo Académico.",
            "Abrir pestaña Trimestres.",
            "Seleccionar cantidad 3 o 4.",
            "Completar fechas de cada trimestre.",
            "Clic Guardar Trimestres.",
            "Activar el trimestre en curso con botón Activar.",
            "Editar fechas con modal si es necesario.",
        ],
        ejemplo="T1 ene–abr · T2 may–ago · T3 sep–dic",
        resultado="Registros en trimesters; actividades podrán asociarse al trimestre activo.",
        validacion="TeacherGradebook/Activity no fallan por 'trimestre inactivo'.",
        errores=[
            ["No guarda", "Fechas inválidas / solape", "Corregir rango"],
            ["Actividades bloqueadas", "Ningún trimestre activo", "ActivarTrimestre"],
        ],
        impacto="Evaluación, gradebook, algunos períodos de prematrícula.",
        buenas=["Un trimestre activo a la vez", "No eliminar sin respaldo"],
        checklist=["Fechas correctas", "Trimestre activo", "Año escolar coherente"],
    )

    procedure(
        doc,
        title="9. Jornada nocturna (Shift Noche)",
        objetivo="Crear la jornada canónica Noche y asociarla a grupos.",
        descripcion="La oferta nocturna y la generación de bloques dependenden del nombre 'Noche'.",
        deps=["School"],
        ruta="/AcademicCatalog/Index → pestaña Jornadas · POST /AcademicCatalog/CreateShift",
        roles="admin / secretaria / director (Authorize)",
        pantalla="Catálogo · Jornadas",
        campos=[["Name", "*", "Debe ser Noche"], ["IsActive", "*", "true"], ["DisplayOrder", "No", "Orden visual"]],
        pasos=[
            "Abrir Catálogo Académico → Jornadas.",
            "Crear jornada con nombre exacto Noche.",
            "Verificar IsActive.",
            "Al crear grupos, asignar ShiftId a Noche.",
        ],
        ejemplo="Name = Noche",
        resultado="Shift disponible; GetOrCreateBySchoolAndNameAsync('Noche') usable por servicios.",
        validacion="Grupos nocturnos muestran jornada Noche; ScheduleConfiguration genera bloques con ese ShiftId.",
        errores=[["Grupos sin oferta prematriculación", "Shift distinto de Noche", "Reasignar ShiftId"]],
        impacto="Prematrícula nocturna, TimeSlots, StudentAssignment.ShiftId.",
        buenas=["Un único shift Noche activo", "No renombrar a 'Nocturna' si servicios buscan 'Noche'"],
        checklist=["Existe Noche", "Grupos vinculados"],
    )

    # Catalog brief
    heading(doc, "10. Catálogo académico — Grados, Grupos, Materias, Áreas, Especialidades", 1)
    para(doc, "Todo se administra en /AcademicCatalog/Index (pestañas) o CRUD individuales /GradeLevel, /Group, /Subject, /Area, /Specialty.")
    table(
        doc,
        ["Pestaña", "Entidad", "Notas nocturnas"],
        [
            ["Grados", "GradeLevel", "7, 8, 9, 10, 11, 12 según oferta"],
            ["Grupos", "Group", "Asignar Shift = Noche; MaxCapacity para cupos"],
            ["Materias", "Subject", "Puede vincular AreaId"],
            ["Áreas", "Area", "Puede ser global IsGlobal"],
            ["Especialidades", "Specialty", "I/T/A/E según CELOSAM"],
            ["Jornadas", "Shift", "Noche"],
            ["Trimestres", "Trimester", "Ver procedimiento 8"],
        ],
    )
    marker(doc, "/AcademicCatalog/Index — pestañas Grados y Grupos")
    callout(
        doc,
        "Importante — Subsecciones A-1 / A-2",
        "El documento oficial de horarios Word a veces usa 7-A-1 / 8-A-2. En BD nocturna típica los grupos de pre-media son únicos (7-A, 8-A, 9-A) y bachillerato 10-A1..A4. "
        "Crear subsecciones solo si la institución lo requiere explícitamente.",
    )

    procedure(
        doc,
        title="11. Impartición (SubjectAssignment)",
        objetivo="Definir qué materia se imparte en qué grupo/grado/especialidad/área.",
        descripcion="Núcleo del modelo. Sin impartición no hay docente ni prematriculación nocturna.",
        deps=["Specialty", "Area", "Subject", "GradeLevel", "Group nocturno"],
        ruta="/SubjectAssignment/Index (+ /AcademicCatalog/Upload para carga masiva)",
        roles="admin (menú)",
        pantalla="Catálogo de Asignaciones / SubjectAssignment",
        campos=[
            ["SpecialtyId", "*", "Especialidad"],
            ["AreaId", "*", "Área"],
            ["SubjectId", "*", "Materia"],
            ["GradeLevelId", "*", "Grado"],
            ["GroupId", "*", "Grupo nocturno"],
        ],
        pasos=[
            "Abrir /SubjectAssignment/Index.",
            "Crear combinaciones materia-grupo requeridas.",
            "Opcional: carga masiva AcademicCatalog/Upload.",
            "Verificar listado GetAllAssignments.",
        ],
        ejemplo="Español · Grado 10 · Grupo 10-A1 · Especialidad Informática",
        resultado="Registros subject_assignments listos para TeacherAssignment.",
        validacion="En prematriculación, el grado aparece en oferta nocturna.",
        errores=[["No aparece grado en Create prematriculación", "Sin SA en grupo Noche", "Crear SubjectAssignment"]],
        impacto="Desbloquea docentes, horarios y oferta académica.",
        buenas=["Nombrar grupos consistente", "Cerrar status Closed solo cuando aplique"],
        checklist=["SA clave por grupo", "Grupos Noche"],
    )

    procedure(
        doc,
        title="12. Profesores — alta y TeacherAssignment",
        objetivo="Registrar docentes y vincularlos a imparticiones.",
        descripcion="Usuario role=teacher; luego TeacherAssignment une TeacherId + SubjectAssignmentId.",
        deps=["SubjectAssignment", "Escuela"],
        ruta="/User/Index · /TeacherAssignment/Index · /AcademicAssignment/Upload",
        roles="admin",
        pantalla="Usuarios + Asignar Docentes",
        campos=[
            ["Name / LastName / Email / Password", "*", "Alta usuario"],
            ["Role", "*", "teacher"],
            ["SubjectAssignment", "*", "En TeacherAssignment"],
        ],
        pasos=[
            "Crear usuario teacher en /User/Index.",
            "Ir a Asignar Docentes.",
            "Seleccionar docente y marcar imparticiones.",
            "Guardar (SaveAssignments).",
            "Verificar cascadas GetAreas/Subjects/Grades/Groups.",
        ],
        ejemplo="Max Agames → Español en 10-A1..12-A4",
        resultado="teacher_assignments creados; dropdown de ByTeacher se llena.",
        validacion="En /Schedule/ByTeacher el docente muestra opciones de materia-grupo.",
        errores=[["Sin opciones en celda horario", "Sin TeacherAssignment", "Asignar imparticiones"],
                 ["Email no único", "Constraint", "Otro email"]],
        impacto="Horarios y gradebook.",
        buenas=["Un correo institucional por docente", "Carga masiva AcademicAssignment/Upload para volumen"],
        checklist=["Usuario teacher", "≥1 TeacherAssignment", "Visible en ByTeacher"],
    )

    procedure(
        doc,
        title="13. Estudiantes, acudientes y matrícula operativa",
        objetivo="Dejar estudiantes matriculados en grupo/grado/año para evaluación y asistencia.",
        descripcion="Matrícula = StudentAssignment (EnrollmentType Nocturno por defecto). Alternativa: flujo prematriculación.",
        deps=["GradeLevel", "Group Noche", "AcademicYear", "User estudiante"],
        ruta="/User/Index · /StudentAssignment/Index · /StudentAssignment/Upload",
        roles="admin; director/secretaria según Authorize de acciones",
        pantalla="Asignar Estudiantes",
        campos=[
            ["Usuario estudiante", "*", "role estudiante/student"],
            ["GradeLevelId / GroupId", "*", "Grado y grupo"],
            ["AcademicYearId", "*", "Año activo"],
            ["ShiftId", "Recomendado", "Heredado del grupo Noche"],
            ["Acudiente ParentId", "Menores", "Rol acudiente vinculado"],
        ],
        pasos=[
            "Crear usuario estudiante.",
            "Opcional: crear acudiente y vincular.",
            "Asignar en StudentAssignment (individual o upload).",
            "Verificar IsActive.",
            "Para cambio de grupo/grado usar UpdateGroupAndGrade.",
        ],
        ejemplo="Estudiante 10-A1 Noche 2026",
        resultado="StudentAssignment activo; puede recibir notas/asistencia.",
        validacion="Gradebook permite cargar notas del grupo.",
        errores=[["No guarda notas", "Sin StudentAssignment activa", "Crear asignación"],
                 ["Menor sin acudiente en prematrícula", "Regla de negocio", "Crear/vincular acudiente"]],
        impacto="Toda operación académica del alumno.",
        buenas=["Usar prematrícula cuando el proceso institucional lo exige (pago)", "Mantener un assignment activo por año"],
        checklist=["Usuario creado", "Assignment activo", "Grupo Noche"],
    )

    procedure(
        doc,
        title="14. Prematrícula (período, solicitud, pago, confirmación)",
        objetivo="Habilitar el portal de prematrícula y convertir solicitudes pagadas en matrícula.",
        descripcion="Período define ventana y cupos. Create valida ofertas nocturnas. ConfirmMatriculation crea StudentAssignment.",
        deps=["SubjectAssignments nocturnos", "AcademicYear", "Trimester (UI)", "Grupos con cupo"],
        ruta="/PrematriculationPeriod · /Prematriculation/Create · /Prematriculation/ConfirmMatriculation · /Payment",
        roles="admin/superadmin período; acudiente/student create; contable/admin pagos",
        pantalla="Período + Nueva Prematrícula + Pagos",
        campos=[
            ["Fecha inicio/fin período", "*", "Ventana"],
            ["MaxCapacityPerGroup", "*", "Cupo"],
            ["AutoAssignByShift", "No", "Asignación automática"],
            ["MaxSubjectsAllowed", "Modular", "Tope materias"],
            ["AcademicYearId / TrimesterId", "*", "En formulario período"],
        ],
        pasos=[
            "Crear período activo en /PrematriculationPeriod/Create (URL directa si no está en menú).",
            "Verificar que grupos nocturnos tienen SubjectAssignments.",
            "Acudiente/estudiante: /Prematriculation/Create.",
            "Registrar pago (portal o contabilidad).",
            "Confirmar matrícula (automática al pago confirmado o ConfirmMatriculation).",
            "Si modular: ModularSubjects → FinalizeModular.",
        ],
        ejemplo="Período enero 2026 · cupo 35 · AutoAssignByShift=true",
        resultado="Estado Matriculado + StudentAssignment.",
        validacion="Estudiante aparece en StudentAssignment; portal estudiante operativo.",
        errores=[
            ["Período no disponible", "Fuera de fechas / inactivo", "Corregir período"],
            ["Sin grados", "Sin SA Noche", "Completar catálogo"],
            ["Cupo lleno", "MaxCapacity", "Otro grupo / ampliar cupo"],
        ],
        impacto="Ingreso institucional controlado por pagos.",
        buenas=["Abrir período solo tras catálogo y docentes listos"],
        checklist=["Período activo", "Oferta Noche", "Pago→Matriculado"],
    )

    # HORARIOS — very detailed
    heading(doc, "15. Horarios — Sección crítica", 1)
    para(
        doc,
        "El módulo de horarios combina tres pantallas: Configuración de jornada (genera bloques Noche), "
        "Ajuste de TimeSlots, y Carga por docente (ScheduleEntry).",
        bold=True,
        color=NAVY,
    )

    procedure(
        doc,
        title="15.1 Configurar jornada nocturna y generar bloques",
        objetivo="Crear automáticamente los TimeSlots de Noche según hora de inicio, duración y cantidad.",
        descripcion="Vista solo nocturna. Al guardar puede reemplazar bloques previos. Si hay schedule_entries puede exigir force/limpieza.",
        deps=["School", "Shift Noche (se crea si falta)"],
        ruta="/ScheduleConfiguration/Index · POST SaveConfiguration",
        roles="admin, director",
        pantalla="Configuración de jornada nocturna",
        campos=[
            ["NightStartTime", "*", "Ej. 18:00 o 18:10 (24h)"],
            ["NightBlockDurationMinutes", "*", "Ej. 50"],
            ["NightBlockCount", "*", "Ej. 6"],
            ["Recreo / RecessAfter…", "Opcional", "Según UI de recreo"],
        ],
        pasos=[
            "Menú Horarios → Configuración de jornada.",
            "Confirmar mensaje 'solo jornada nocturna'.",
            "Completar hora inicio, duración, cantidad.",
            "Revisar cálculo estimado en pantalla.",
            "Guardar. Leer TempData Success/Error.",
            "Ir a TimeSlot/Manage y verificar Bloque 1..N con Shift Noche.",
        ],
        ejemplo="18:10 · 50 min · 6 bloques → 18:10-19:00 … 22:20-23:10",
        resultado="time_slots activos con Shift Noche; mensaje de bloques generados.",
        validacion="Lista TimeSlot muestra 6 bloques nocturnos; no deben prevalecer bloques 07:00 legacy.",
        errores=[
            ["No regenera", "Hay schedule_entries", "Eliminar entradas o force según servicio"],
            ["Bloques mañana mezclados", "Bootstrap default", "Regenerar con ScheduleConfiguration o DeleteAll"],
        ],
        impacto="Base de la grilla ByTeacher.",
        buenas=["Fijar horario oficial antes de cargar celdas", "Respaldar schedule_entries antes de regenerar"],
        checklist=["Solo Noche", "Horas correctas", "Sin basura mañana"],
    )

    procedure(
        doc,
        title="15.2 Ajustar / eliminar bloques (TimeSlot)",
        objetivo="Afinar nombres, horas o limpiar bloques.",
        descripcion="Manage es la entrada de menú. DeleteAll exige escribir ELIMINAR y borra primero schedule_entries.",
        deps=["School"],
        ruta="/TimeSlot/Manage · /TimeSlot/Index · DeleteAll",
        roles="admin, director",
        pantalla="Ajustar bloques horarios",
        campos=[["Name", "*", "Bloque 1…"], ["StartTime/EndTime", "*", "time"], ["ShiftId", "Recomendado", "Noche"], ["IsActive", "*", "true"]],
        pasos=[
            "Abrir Ajustar bloques horarios.",
            "Editar si una hora oficial difiere levemente.",
            "Para limpieza total: Index → Eliminar todos → confirmar ELIMINAR.",
            "Volver a ScheduleConfiguration para regenerar.",
        ],
        ejemplo="Renombrar Bloque 1 · 18:10-19:00",
        resultado="Bloques coherentes con documento oficial de horarios.",
        validacion="ByTeacher muestra columnas de bloque correctas.",
        errores=[["No borra", "Confirmación distinta de ELIMINAR", "Escribir exacto"],
                 ["FK error", "Orden delete", "Usar DeleteAll del sistema"]],
        impacto="Cambia toda la grilla.",
        buenas=["No editar a medias con horarios ya publicados sin avisar docentes"],
        checklist=["Bloques Noche", "Orden DisplayOrder"],
    )

    procedure(
        doc,
        title="15.3 Cargar horario por docente (ScheduleEntry)",
        objetivo="Asignar materia-grupo a cada día/bloque del docente.",
        descripcion="Celda a celda. Cada SaveEntry crea ScheduleEntry ligado a TeacherAssignment.",
        deps=["TeacherAssignment", "TimeSlots Noche", "AcademicYear activo"],
        ruta="/Schedule/ByTeacher · POST /Schedule/SaveEntry · POST /Schedule/DeleteEntry",
        roles="admin, director, teacher (solo propias asignaciones)",
        pantalla="Horario por Docente",
        campos=[
            ["TeacherId", "*", "Admin/director eligen"],
            ["AcademicYearId", "*", "Año canónico"],
            ["TeacherAssignmentId", "*", "Materia-grupo"],
            ["TimeSlotId", "*", "Bloque"],
            ["DayOfWeek", "*", "1=Lunes … 5=Viernes (soporta 1–7)"],
        ],
        pasos=[
            "Menú Horarios → Cargar horarios por docente.",
            "Seleccionar docente y año → Cargar horario.",
            "En cada celda elegir la impartición correcta.",
            "El sistema guarda vía AJAX SaveEntry.",
            "Para quitar: eliminar celda (DeleteEntry).",
            "Docente: puede cargar el suyo si IsEditable.",
        ],
        ejemplo="Lunes 1ra → Español 11-A1 (Max Agames)",
        resultado="schedule_entries; estudiante verá en MySchedule tras matrícula/SSA.",
        validacion="ListJsonByTeacher retorna entradas; sin alertas de conflicto.",
        errores=[
            ["Conflicto docente", "Traslape horario mismo día", "Mover otra materia"],
            ["Conflicto grupo", "Grupo ya ocupado en bloque", "Revisar otro docente"],
            ["Sin año", "HasNoAcademicYears", "Crear/activar año"],
            ["Dropdown vacío", "Sin TeacherAssignment", "Asignar docente"],
        ],
        impacto="Operación semanal de clases; vista estudiante.",
        buenas=[
            "Cargar desde documento oficial de horarios",
            "Un bloque = una impartición",
            "Resolver códigos ambiguos I-T-A-E antes de guardar",
            "No existe importador nativo Word/Excel — usar carga manual o scripts autorizados",
        ],
        checklist=["Año canónico", "Bloques Noche", "TA completos", "Sin conflictos", "Muestra en ByTeacher"],
    )

    heading(doc, "15.4 Reglas, restricciones y casos especiales de horarios", 2)
    bullet(doc, "Un docente no puede tener dos clases con horas traslapadas el mismo día.")
    bullet(doc, "Un grupo no puede tener dos clases traslapadas el mismo día.")
    bullet(doc, "Teacher solo edita sus TeacherAssignments.")
    bullet(doc, "Room/Salón: no operativo como módulo de asignación de aulas.")
    bullet(doc, "Recreos: generables desde configuración si se habilita en UI.")
    bullet(doc, "Casos especiales CELOSAM: codigos I-T-A-E del papel deben mapearse a un grupo real (10-A1...).")
    marker(doc, "/Schedule/ByTeacher con grilla llena de ejemplo")
    doc.add_page_break()

    heading(doc, "16. Evaluación, actividades y asistencia", 1)
    procedure(
        doc,
        title="16.1 Actividades y Gradebook",
        objetivo="Registrar actividades y calificaciones del trimestre activo.",
        descripcion="ActivityController + TeacherGradebook. Requiere trimestre activo y StudentAssignment.",
        deps=["Trimester activo", "TeacherAssignment", "StudentAssignment"],
        ruta="/Activity/* · /TeacherGradebook/Index",
        roles="teacher/docente (+ admin/director en Activity)",
        pantalla="Portal Docente / Gradebook",
        campos=[["Trimestre", "*", "Activo"], ["Grupo/Materia", "*", "Desde TA"], ["Puntaje", "*", "Según actividad"]],
        pasos=["Activar trimestre", "Crear actividad", "Abrir Gradebook", "Guardar scores SaveScores"],
        ejemplo="Examen T1 Matemática 10-A1",
        resultado="StudentActivityScore persistidos.",
        validacion="Consultas de promedio / reportes muestran nota.",
        errores=[["Sin asignación estudiante", "Falta StudentAssignment", "Matricular"],
                 ["Trimestre inactivo", "ValidateTrimesterActive", "Activar trimestre"]],
        impacto="Promoción, reportes, reglas de prematrícula por reprobadas.",
        buenas=["Cerrar trimestre formalmente"],
        checklist=["Trimestre activo", "Notas de prueba cargadas"],
    )

    procedure(
        doc,
        title="16.2 Asistencia",
        objetivo="Registrar asistencia del grupo.",
        descripcion="AttendanceController con roles admin,secretaria,teacher,docente,director.",
        deps=["StudentAssignment / estudiantes en grupo"],
        ruta="/Attendance/Index",
        roles="admin,secretaria,teacher,docente,director",
        pantalla="Asistencia",
        campos=[["Grupo/Fecha", "*", "Contexto"], ["Estado asistencia", "*", "Por estudiante"]],
        pasos=["Abrir Attendance", "Seleccionar grupo/fecha", "Marcar y guardar"],
        ejemplo="Asistencia lunes grupo 10-A1",
        resultado="Registros de asistencia.",
        validacion="Reporte/consulta refleja marcas.",
        errores=[["Lista vacía", "Sin estudiantes asignados", "StudentAssignment"]],
        impacto="Seguimiento operativo.",
        buenas=["Registrar el mismo día"],
        checklist=["Grupo con alumnos", "Marca de prueba"],
    )

    # Remaining high-level chapters
    heading(doc, "17. Reportes, CELOSAM, carnets y portales", 1)
    table(
        doc,
        ["Módulo", "Ruta", "Rol típico"],
        [
            ["Reportes estudiante", "/StudentReport", "estudiante"],
            ["Orientación estudiante", "/StudentOrientation", "estudiante"],
            ["Horario estudiante", "/StudentSchedule/MySchedule", "estudiante"],
            ["Portal acudiente académico", "/ParentAcademic", "acudiente/parent"],
            ["Documentos CELOSAM", "/Celosan/Documents", "admin/secretaria"],
            ["Créditos / Convalidaciones", "/Celosan/BulkCredits", "admin/secretaria"],
            ["Equivalencias", "/SuperAdmin/Equivalencies", "admin/secretaria"],
            ["Malla y prerrequisitos", "/SuperAdmin/CurriculumTracks", "admin/secretaria"],
            ["Reportes CELOSAM", "/Celosan/Reports", "admin/secretaria"],
            ["Carnet estudiantil", "/StudentIdCard/ui", "admin/secretaria/superadmin"],
            ["Pagos", "/Payment / PaymentConcept", "contable/admin/acudiente"],
            ["Mensajería", "/Messaging", "según rol"],
            ["Director", "/Director/Index", "director"],
        ],
    )
    marker(doc, "Portal Director /Home o /Director/Index")

    heading(doc, "18. Validación final y certificación operativa (Go-Live)", 1)
    para(doc, "Checklist de certificación — marcar todos antes de abrir a producción masiva.", bold=True)
    for c in [
        "Escuela y admin operativos",
        "Un AcademicYear activo canónico",
        "Shift Noche creado",
        "Grados y grupos nocturnos con cupo",
        "Áreas, especialidades y materias cargadas",
        "Trimestres creados y uno activo",
        "SubjectAssignments completas para oferta",
        "Docentes con TeacherAssignments",
        "TimeSlots solo Noche coherentes con horario oficial",
        "ScheduleEntries de prueba sin conflictos",
        "Estudiante de prueba matriculado (StudentAssignment)",
        "Gradebook guarda nota de prueba",
        "Asistencia marca de prueba",
        "Período de prematrícula (si aplica) probado ponta a ponta",
        "Portales student/acudiente/director smoke-test",
        "Backup BD realizado",
    ]:
        bullet(doc, f"☐ {c}")

    callout(
        doc,
        "Criterio Go-Live",
        "La institución está certificada cuando un estudiante nocturno de punta a punta puede: prematricularse o estar asignado, "
        "ver horario, recibir nota y asistencia, y el admin puede emitir reportes CELOSAM/carnet según alcance contratado.",
        fill="E8F8EF",
        border="1B7A3D",
    )
    doc.add_page_break()

    heading(doc, "19. Matrices y diagramas", 1)
    heading(doc, "19.1 Matriz de dependencias (resumen)", 2)
    table(
        doc,
        ["Entidad", "Depende de", "Habilita"],
        [
            ["Shift Noche", "School", "Groups, TimeSlots"],
            ["SubjectAssignment", "5 FK catálogo", "TA, prematrícula"],
            ["TeacherAssignment", "User+SA", "Horario, notas"],
            ["TimeSlot", "School+Shift", "ScheduleEntry"],
            ["ScheduleEntry", "TA+TS+Year", "MySchedule"],
            ["StudentAssignment", "User+GL+Group+Year", "Notas, asistencia"],
            ["Activity/Score", "Trimester+SA+STA", "Reportes"],
        ],
    )
    heading(doc, "19.2 Diagrama de flujo de configuración (texto)", 2)
    para(
        doc,
        "[SuperAdmin: Escuela] → [Admin login] → [Catálogo: Noche+GL+Grupos+Materias+Trimestres] → "
        "[SubjectAssignment] → [Users] → [TeacherAssignment] → [ScheduleConfiguration] → [ByTeacher] → "
        "[PrematriculationPeriod?] → [Students/Prematrícula] → [Ops Gradebook/Attendance] → [Validación Go-Live]",
        size=10,
    )

    heading(doc, "20. Glosario y limitaciones conocidas", 1)
    table(
        doc,
        ["Término", "Significado en sistema"],
        [
            ["SubjectAssignment", "Impartición materia-grupo-grado-área-especialidad"],
            ["TeacherAssignment", "Docente ↔ impartición"],
            ["StudentAssignment", "Matrícula operativa del estudiante"],
            ["ScheduleEntry", "Celda de horario"],
            ["TimeSlot", "Bloque horario"],
            ["Shift Noche", "Jornada canónica nocturna"],
            ["EnrollmentType Nocturno", "Tipo por defecto de matrícula primaria"],
        ],
    )
    para(doc, "Limitaciones (no inventar trabajo futuro como hecho):", bold=True, color=ORANGE)
    bullet(doc, "Sin UI Index de AcademicYear.")
    bullet(doc, "Sin importador nativo de horarios Word/Excel.")
    bullet(doc, "Salones (Room) no operativos en horario.")
    bullet(doc, "Escala de notas no es un módulo independiente de configuración.")
    bullet(doc, "Algunos menús de prematrícula admin pueden estar ocultos — usar URL directa autorizada.")
    bullet(doc, "TimeSlot Create puede mostrar defaults 07:00 (legacy); preferir ScheduleConfiguration nocturna.")

    doc.add_page_break()
    heading(doc, "Anexo A — Referencia rápida de rutas", 1)
    table(
        doc,
        ["Función", "Ruta"],
        [
            ["Login", "/Auth/Login"],
            ["Dashboard", "/Home/Index"],
            ["Crear escuela", "/SuperAdmin/CreateSchoolWithAdmin"],
            ["Catálogo", "/AcademicCatalog/Index"],
            ["Imparticiones", "/SubjectAssignment/Index"],
            ["Usuarios", "/User/Index"],
            ["Docentes", "/TeacherAssignment/Index"],
            ["Estudiantes", "/StudentAssignment/Index"],
            ["Jornada noche", "/ScheduleConfiguration/Index"],
            ["Bloques", "/TimeSlot/Manage"],
            ["Horario docente", "/Schedule/ByTeacher"],
            ["Horario alumno", "/StudentSchedule/MySchedule"],
            ["Período prematrícula", "/PrematriculationPeriod/Index"],
            ["Nueva prematrícula", "/Prematriculation/Create"],
            ["Gradebook", "/TeacherGradebook/Index"],
            ["Asistencia", "/Attendance/Index"],
            ["Análisis técnico", "Documentacion/EduplanerNocturna/ANALISIS_COMPLETO_EDUPLANER_NOCTURNA.md"],
        ],
    )

    para(
        doc,
        "Fin del Manual Oficial Premium de Configuración Inicial — Eduplaner Nocturna. "
        "Documento generado tras inventario de Controllers/Services/Models/Views/Menús. "
        "Actualizar versiones cuando cambie el código de configuración crítica.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        italic=True,
        size=10,
    )

    doc.save(OUT)
    print(f"OK -> {OUT}")
    print(f"Size KB: {OUT.stat().st_size // 1024}")


if __name__ == "__main__":
    build()
